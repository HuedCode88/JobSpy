import csv
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo
from jobspy import scrape_jobs

 
COLUMN_WIDTHS = {
    "title": 34,
    "company": 28,
    "description": 60,
    "job_url": 52,
    "job_url_direct": 52,
    "location": 24,
}


search_terms = [
    "software engineer New Grad",
    "software engineer I",
    "software engineer 1",
    "software engineer intern",
    "software engineer co-op",
]

results = []

for term in search_terms:
    jobs = scrape_jobs(
        site_name=["indeed", "linkedin", "zip_recruiter", "glassdoor"],
        search_term=term,
        location="Miami, FL",
        results_wanted=20,
        hours_old=24,
        country_indeed="USA",
        fetch_description=True,
        verbose=1,
    )

    results.append(jobs)

all_jobs = pd.concat(results, ignore_index=True)

# Remove the same listing found by multiple searches
all_jobs = all_jobs.drop_duplicates(subset=["job_url"])

print(f"Found {len(all_jobs)} unique jobs")
print(all_jobs.head())

all_jobs.to_csv(
    "jobs.csv",
    quoting=csv.QUOTE_NONNUMERIC,
    escapechar="\\",
    index=False,
)

all_jobs.to_excel("jobs.xlsx", index=False, engine="openpyxl")

workbook = load_workbook("jobs.xlsx")
worksheet = workbook.active
worksheet.title = "Job Results"
worksheet.freeze_panes = "A2"

header_fill = PatternFill("solid", fgColor="1F4E78")
header_font = Font(color="FFFFFF", bold=True)

for cell in worksheet[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

for column_cells in worksheet.columns:
    column_name = column_cells[0].value
    letter = column_cells[0].column_letter
    worksheet.column_dimensions[letter].width = COLUMN_WIDTHS.get(
        column_name, 16
    )

    for cell in column_cells[1:]:
        cell.alignment = Alignment(vertical="top", wrap_text=True)

        if column_name in {"job_url", "job_url_direct"} and cell.value:
            cell.hyperlink = cell.value
            cell.font = Font(color="0563C1", underline="single")

for row in worksheet.iter_rows(min_row=2):
    worksheet.row_dimensions[row[0].row].height = 48

if worksheet.max_row > 1 and worksheet.max_column > 0:
    table = Table(displayName="JobResults", ref=worksheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(table)

workbook.save("jobs.xlsx")