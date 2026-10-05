from __future__ import annotations

import argparse
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


COLUMN_WIDTHS = {
    "title": 34,
    "company": 28,
    "description": 60,
    "job_url": 52,
    "job_url_direct": 52,
    "location": 24,
}


def format_jobs(input_path: Path, output_path: Path) -> None:
    if not input_path.exists():
        raise FileNotFoundError(f"Input file was not found: {input_path}")

    jobs = pd.read_csv(input_path)
    jobs.to_excel(output_path, index=False, engine="openpyxl")

    workbook = load_workbook(output_path)
    worksheet = workbook.active
    worksheet.title = "Job Results"
    worksheet.freeze_panes = "A2"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for column_cells in worksheet.columns:
        column_name = column_cells[0].value
        letter = column_cells[0].column_letter
        worksheet.column_dimensions[letter].width = COLUMN_WIDTHS.get(
            column_name, 16
        )

        for cell in column_cells[1:]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

            if column_name in {"job_url", "job_url_direct"} and cell.value:
                cell.hyperlink = cell.value
                cell.font = Font(color="0563C1", underline="single")

    for row in worksheet.iter_rows(min_row=2):
        worksheet.row_dimensions[row[0].row].height = 48

    if worksheet.max_row > 1 and worksheet.max_column > 0:
        table = Table(displayName="JobResults", ref=worksheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        worksheet.add_table(table)

    workbook.save(output_path)
    print(f"Formatted {len(jobs)} jobs from {input_path} into {output_path}")


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Format an existing JobSpy CSV without making web requests."
    )
    parser.add_argument(
        "input_csv",
        nargs="?",
        default="jobs.csv",
        help="Existing JobSpy CSV file (default: jobs.csv)",
    )
    parser.add_argument(
        "output_xlsx",
        nargs="?",
        default="jobs_readable.xlsx",
        help="Formatted Excel output file (default: jobs_readable.xlsx)",
    )
    args = parser.parse_args()

    format_jobs(Path(args.input_csv), Path(args.output_xlsx))


if __name__ == "__main__":
    main()
